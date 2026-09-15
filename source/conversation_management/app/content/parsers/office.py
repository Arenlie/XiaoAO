from __future__ import annotations

import asyncio
import io
from pathlib import Path
from typing import Any

from app.attachments.contracts import AttachmentDescriptor
from app.content.contracts import ContentKind, ContentUnderstandingResult, ExtractionStatus


class PdfContentParser:
    name = "pdf.v1"
    supported_kinds = {"pdf"}

    async def parse(self, *, descriptor: AttachmentDescriptor, data: bytes, max_characters: int) -> ContentUnderstandingResult:
        def _parse() -> tuple[str, list[dict[str, Any]], list[str]]:
            warnings: list[str] = []
            try:
                from pypdf import PdfReader
            except ImportError:
                return "", [], ["未安装 pypdf，仅能通过多模态模型原生文件能力分析 PDF。"]
            reader = PdfReader(io.BytesIO(data))
            chunks: list[str] = []
            pages: list[dict[str, Any]] = []
            used = 0
            for index, page in enumerate(reader.pages, start=1):
                try:
                    text = page.extract_text() or ""
                except Exception:
                    text = ""
                    warnings.append(f"第 {index} 页文本提取失败。")
                preview = text[:2000]
                pages.append({"page_number": index, "characters": len(text), "preview": preview})
                if used < max_characters:
                    remaining = max_characters - used
                    chunks.append(f"\n[第{index}页]\n{text[:remaining]}")
                    used += min(len(text), remaining)
            return "".join(chunks).strip(), pages, warnings

        extracted, pages, warnings = await asyncio.to_thread(_parse)
        status = ExtractionStatus.COMPLETED if extracted else ExtractionStatus.PARTIAL
        return ContentUnderstandingResult(
            content_id=str(descriptor.attachment_id),
            attachment_id=descriptor.attachment_id,
            kind=ContentKind.PDF,
            extraction_status=status,
            summary=f"PDF 文件 {descriptor.filename}，共 {len(pages) or '未知'} 页。",
            extracted_text=extracted,
            pages=pages,
            capabilities_required=["pdf_understanding"],
            warnings=warnings,
            parser_name=self.name,
        )


class DocumentContentParser:
    name = "document.v1"
    supported_kinds = {"document"}

    async def parse(self, *, descriptor: AttachmentDescriptor, data: bytes, max_characters: int) -> ContentUnderstandingResult:
        extension = Path(descriptor.filename).suffix.lower()
        if extension != ".docx":
            return ContentUnderstandingResult(
                content_id=str(descriptor.attachment_id),
                attachment_id=descriptor.attachment_id,
                kind=ContentKind.DOCUMENT,
                extraction_status=ExtractionStatus.PARTIAL,
                summary=f"文档文件 {descriptor.filename}。",
                capabilities_required=["document_understanding"],
                warnings=["当前结构化解析仅完整支持 DOCX；该格式将优先交给模型原生文件能力。"],
                parser_name=self.name,
            )

        def _parse() -> tuple[str, list[dict[str, Any]], list[str]]:
            try:
                from docx import Document
            except ImportError:
                return "", [], ["未安装 python-docx。"]
            document = Document(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            tables: list[dict[str, Any]] = []
            for index, table in enumerate(document.tables, start=1):
                rows = [[cell.text for cell in row.cells] for row in table.rows[:30]]
                tables.append({"index": index, "rows": len(table.rows), "columns": len(table.columns), "preview_rows": rows})
            return "\n".join(paragraphs)[:max_characters], tables, []

        text, tables, warnings = await asyncio.to_thread(_parse)
        return ContentUnderstandingResult(
            content_id=str(descriptor.attachment_id),
            attachment_id=descriptor.attachment_id,
            kind=ContentKind.DOCUMENT,
            extraction_status=ExtractionStatus.COMPLETED if text or tables else ExtractionStatus.PARTIAL,
            summary=f"Word 文档 {descriptor.filename}，提取到正文和 {len(tables)} 个表格。",
            extracted_text=text,
            tables=tables,
            capabilities_required=["document_understanding"],
            warnings=warnings,
            parser_name=self.name,
        )


class SpreadsheetContentParser:
    name = "spreadsheet.v1"
    supported_kinds = {"spreadsheet"}

    async def parse(self, *, descriptor: AttachmentDescriptor, data: bytes, max_characters: int) -> ContentUnderstandingResult:
        extension = Path(descriptor.filename).suffix.lower()
        if extension in {".csv", ".tsv"}:
            from app.content.parsers.text import CsvContentParser
            return await CsvContentParser().parse(
                descriptor=descriptor, data=data, max_characters=max_characters
            )
        if extension != ".xlsx":
            return ContentUnderstandingResult(
                content_id=str(descriptor.attachment_id),
                attachment_id=descriptor.attachment_id,
                kind=ContentKind.SPREADSHEET,
                extraction_status=ExtractionStatus.PARTIAL,
                summary=f"表格文件 {descriptor.filename}。",
                capabilities_required=["spreadsheet_inspection"],
                warnings=["当前结构化解析完整支持 XLSX/CSV/TSV；XLS 格式将交给模型原生文件能力或专用工具。"],
                parser_name=self.name,
            )

        def _parse() -> tuple[list[dict[str, Any]], str, list[str]]:
            try:
                from openpyxl import load_workbook
            except ImportError:
                return [], "", ["未安装 openpyxl。"]
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            sheets: list[dict[str, Any]] = []
            text_parts: list[str] = []
            warnings: list[str] = []
            used = 0
            for sheet in workbook.worksheets:
                preview_rows: list[list[Any]] = []
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    values = [value for value in row]
                    if row_index <= 20:
                        preview_rows.append(values[:50])
                    if used < max_characters and row_index <= 200:
                        line = "\t".join("" if value is None else str(value) for value in values[:50])
                        remaining = max_characters - used
                        text_parts.append(f"[{sheet.title}!{row_index}] {line[:remaining]}")
                        used += min(len(line), remaining)
                    if row_index >= 2000:
                        warnings.append(f"工作表 {sheet.title} 较大，仅扫描前 2000 行进行结构理解。")
                        break
                headers = preview_rows[0] if preview_rows else []
                sheets.append({
                    "name": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    "headers": headers[:50],
                    "preview_rows": preview_rows[:10],
                })
            workbook.close()
            return sheets, "\n".join(text_parts), warnings

        sheets, text, warnings = await asyncio.to_thread(_parse)
        return ContentUnderstandingResult(
            content_id=str(descriptor.attachment_id),
            attachment_id=descriptor.attachment_id,
            kind=ContentKind.SPREADSHEET,
            extraction_status=ExtractionStatus.COMPLETED if sheets else ExtractionStatus.PARTIAL,
            summary=f"Excel 文件 {descriptor.filename}，包含 {len(sheets)} 个工作表。",
            extracted_text=text[:max_characters],
            sheets=sheets,
            capabilities_required=["spreadsheet_inspection"],
            warnings=warnings,
            parser_name=self.name,
        )


class PresentationContentParser:
    name = "presentation.v1"
    supported_kinds = {"presentation"}

    async def parse(self, *, descriptor: AttachmentDescriptor, data: bytes, max_characters: int) -> ContentUnderstandingResult:
        extension = Path(descriptor.filename).suffix.lower()
        if extension != ".pptx":
            return ContentUnderstandingResult(
                content_id=str(descriptor.attachment_id),
                attachment_id=descriptor.attachment_id,
                kind=ContentKind.PRESENTATION,
                extraction_status=ExtractionStatus.PARTIAL,
                summary=f"演示文稿 {descriptor.filename}。",
                capabilities_required=["presentation_understanding"],
                warnings=["当前结构化解析完整支持 PPTX；PPT 格式将交给模型原生文件能力或专用工具。"],
                parser_name=self.name,
            )

        def _parse() -> tuple[list[dict[str, Any]], str, list[str]]:
            try:
                from pptx import Presentation
            except ImportError:
                return [], "", ["未安装 python-pptx。"]
            presentation = Presentation(io.BytesIO(data))
            slides: list[dict[str, Any]] = []
            chunks: list[str] = []
            for index, slide in enumerate(presentation.slides, start=1):
                texts: list[str] = []
                for shape in slide.shapes:
                    text = getattr(shape, "text", "")
                    if text and str(text).strip():
                        texts.append(str(text).strip())
                joined = "\n".join(texts)
                slides.append({"slide_number": index, "text": joined[:3000]})
                chunks.append(f"[第{index}页]\n{joined}")
            return slides, "\n".join(chunks)[:max_characters], []

        slides, text, warnings = await asyncio.to_thread(_parse)
        return ContentUnderstandingResult(
            content_id=str(descriptor.attachment_id),
            attachment_id=descriptor.attachment_id,
            kind=ContentKind.PRESENTATION,
            extraction_status=ExtractionStatus.COMPLETED if slides else ExtractionStatus.PARTIAL,
            summary=f"PPTX 文件 {descriptor.filename}，共 {len(slides)} 页。",
            extracted_text=text,
            slides=slides,
            capabilities_required=["presentation_understanding"],
            warnings=warnings,
            parser_name=self.name,
        )
