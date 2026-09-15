from __future__ import annotations

import asyncio
import csv
import io
from collections import defaultdict
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from app.attachments.contracts import AttachmentDescriptor, AttachmentKind
from app.attachments.service import AttachmentService
from app.content.understanding_service import ContentUnderstandingService
from app.domain.exceptions import AppError
from app.tools.contracts import (
    ToolCallRequest,
    ToolCallResult,
    ToolDescriptor,
    ToolProviderType,
    ToolResultStatus,
)

FILE_SEARCH_TOOL_ID = "file.search"
SPREADSHEET_INSPECT_TOOL_ID = "spreadsheet.inspect"
SPREADSHEET_READ_RANGE_TOOL_ID = "spreadsheet.read_range"
SPREADSHEET_FILTER_TOOL_ID = "spreadsheet.filter"
SPREADSHEET_AGGREGATE_TOOL_ID = "spreadsheet.aggregate"


def default_file_tool_descriptors() -> list[ToolDescriptor]:
    attachment_property = {
        "type": "string",
        "format": "uuid",
        "description": "消息附件的 attachment_id",
    }
    return [
        ToolDescriptor(
            tool_id=FILE_SEARCH_TOOL_ID,
            display_name="文件内容检索",
            provider_type=ToolProviderType.LOCAL,
            description="在已本地解析的附件内容中检索与问题最相关的片段，并保留片段定位。",
            input_schema={
                "type": "object",
                "properties": {
                    "attachment_id": attachment_property,
                    "query": {"type": "string"},
                    "max_characters": {
                        "type": "integer",
                        "minimum": 500,
                        "maximum": 30000,
                        "default": 10000,
                    },
                },
                "required": ["attachment_id", "query"],
                "additionalProperties": False,
            },
            output_schema={"type": "object"},
            supports_attachments=True,
            supported_attachment_kinds=["pdf", "document", "spreadsheet", "text"],
            timeout_seconds=30,
        ),
        ToolDescriptor(
            tool_id=SPREADSHEET_INSPECT_TOOL_ID,
            display_name="表格结构检查",
            provider_type=ToolProviderType.LOCAL,
            description="确定性读取 XLSX、CSV 或 TSV 的工作表、表头、行列数量和少量预览。",
            input_schema={
                "type": "object",
                "properties": {"attachment_id": attachment_property},
                "required": ["attachment_id"],
                "additionalProperties": False,
            },
            output_schema={"type": "object"},
            supports_attachments=True,
            supported_attachment_kinds=["spreadsheet"],
            timeout_seconds=60,
        ),
        ToolDescriptor(
            tool_id=SPREADSHEET_READ_RANGE_TOOL_ID,
            display_name="表格范围读取",
            provider_type=ToolProviderType.LOCAL,
            description="按工作表和单元格范围确定性读取表格数据。",
            input_schema={
                "type": "object",
                "properties": {
                    "attachment_id": attachment_property,
                    "sheet_name": {"type": ["string", "null"]},
                    "cell_range": {
                        "type": "string",
                        "description": "A1:H50 格式；CSV/TSV 使用行列范围同样适用",
                        "default": "A1:Z100",
                    },
                },
                "required": ["attachment_id"],
                "additionalProperties": False,
            },
            output_schema={"type": "object"},
            supports_attachments=True,
            supported_attachment_kinds=["spreadsheet"],
            timeout_seconds=60,
        ),
        ToolDescriptor(
            tool_id=SPREADSHEET_FILTER_TOOL_ID,
            display_name="表格精确筛选",
            provider_type=ToolProviderType.LOCAL,
            description="按列名和条件确定性筛选 XLSX、CSV 或 TSV，并返回限定数量的结果。",
            input_schema={
                "type": "object",
                "properties": {
                    "attachment_id": attachment_property,
                    "sheet_name": {"type": ["string", "null"]},
                    "filters": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "column": {"type": "string"},
                                "operator": {
                                    "type": "string",
                                    "enum": [
                                        "eq", "ne", "contains", "starts_with", "ends_with",
                                        "gt", "gte", "lt", "lte", "in", "is_empty", "not_empty",
                                    ],
                                },
                                "value": {},
                            },
                            "required": ["column", "operator"],
                            "additionalProperties": False,
                        },
                        "default": [],
                    },
                    "select_columns": {
                        "type": "array",
                        "items": {"type": "string"},
                        "default": [],
                    },
                    "limit": {
                        "type": "integer",
                        "minimum": 1,
                        "maximum": 500,
                        "default": 100,
                    },
                },
                "required": ["attachment_id"],
                "additionalProperties": False,
            },
            output_schema={"type": "object"},
            supports_attachments=True,
            supported_attachment_kinds=["spreadsheet"],
            timeout_seconds=120,
        ),
        ToolDescriptor(
            tool_id=SPREADSHEET_AGGREGATE_TOOL_ID,
            display_name="表格精确聚合",
            provider_type=ToolProviderType.LOCAL,
            description="按分组列执行 count、sum、avg、min、max 等确定性统计。",
            input_schema={
                "type": "object",
                "properties": {
                    "attachment_id": attachment_property,
                    "sheet_name": {"type": ["string", "null"]},
                    "filters": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "column": {"type": "string"},
                                "operator": {"type": "string"},
                                "value": {},
                            },
                            "required": ["column", "operator"],
                        },
                        "default": [],
                    },
                    "group_by": {
                        "type": "array",
                        "items": {"type": "string"},
                        "default": [],
                    },
                    "metrics": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "operation": {
                                    "type": "string",
                                    "enum": ["count", "sum", "avg", "min", "max"],
                                },
                                "column": {"type": ["string", "null"]},
                                "alias": {"type": ["string", "null"]},
                            },
                            "required": ["operation"],
                        },
                        "minItems": 1,
                    },
                    "limit": {
                        "type": "integer",
                        "minimum": 1,
                        "maximum": 500,
                        "default": 100,
                    },
                },
                "required": ["attachment_id", "metrics"],
                "additionalProperties": False,
            },
            output_schema={"type": "object"},
            supports_attachments=True,
            supported_attachment_kinds=["spreadsheet"],
            timeout_seconds=180,
        ),
    ]


class FileAndSpreadsheetToolHandlers:
    def __init__(
        self,
        *,
        attachment_service: AttachmentService,
        content_service: ContentUnderstandingService,
        max_scan_rows: int = 100000,
    ) -> None:
        self.attachment_service = attachment_service
        self.content_service = content_service
        self.max_scan_rows = max(1000, max_scan_rows)

    @staticmethod
    def _attachment_id(request: ToolCallRequest) -> UUID:
        raw = request.arguments.get("attachment_id")
        try:
            return UUID(str(raw))
        except (TypeError, ValueError) as exc:
            raise AppError("ATTACHMENT_ID_INVALID", "attachment_id 不是合法 UUID", 422) from exc

    async def file_search(
        self, request: ToolCallRequest, data_access_token: str | None
    ) -> ToolCallResult:
        del data_access_token
        attachment_id = self._attachment_id(request)
        query = str(request.arguments.get("query") or "").strip()
        max_characters = min(
            30000, max(500, int(request.arguments.get("max_characters") or 10000))
        )
        result = await self.content_service.search_attachment(
            attachment_id=attachment_id,
            user_token=request.user_token,
            query=query,
            max_characters=max_characters,
        )
        return ToolCallResult(
            tool_id=FILE_SEARCH_TOOL_ID,
            status=ToolResultStatus.SUCCESS,
            content="\n\n".join(item["content"] for item in result["matches"]),
            structured_content=result,
            metadata={"deterministic": True},
        )

    async def spreadsheet_inspect(
        self, request: ToolCallRequest, data_access_token: str | None
    ) -> ToolCallResult:
        del data_access_token
        descriptor, data = await self.attachment_service.read_owned(
            attachment_id=self._attachment_id(request), user_token=request.user_token
        )
        self._require_spreadsheet(descriptor)
        result = await asyncio.to_thread(self._inspect_sync, descriptor, data)
        return ToolCallResult(
            tool_id=SPREADSHEET_INSPECT_TOOL_ID,
            status=ToolResultStatus.SUCCESS,
            content=result,
            structured_content=result,
            metadata={"deterministic": True},
        )

    async def spreadsheet_read_range(
        self, request: ToolCallRequest, data_access_token: str | None
    ) -> ToolCallResult:
        del data_access_token
        descriptor, data = await self.attachment_service.read_owned(
            attachment_id=self._attachment_id(request), user_token=request.user_token
        )
        self._require_spreadsheet(descriptor)
        cell_range = str(request.arguments.get("cell_range") or "A1:Z100").upper()
        try:
            min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        except ValueError as exc:
            raise AppError("CELL_RANGE_INVALID", "cell_range 必须使用 A1:H50 格式", 422) from exc
        if max_row - min_row + 1 > 1000 or max_col - min_col + 1 > 100:
            raise AppError("CELL_RANGE_TOO_LARGE", "单次最多读取 1000 行、100 列", 422)
        result = await asyncio.to_thread(
            self._read_range_sync,
            descriptor,
            data,
            request.arguments.get("sheet_name"),
            min_row,
            max_row,
            min_col,
            max_col,
        )
        return ToolCallResult(
            tool_id=SPREADSHEET_READ_RANGE_TOOL_ID,
            status=ToolResultStatus.SUCCESS,
            content=result,
            structured_content=result,
            metadata={"deterministic": True},
        )

    async def spreadsheet_filter(
        self, request: ToolCallRequest, data_access_token: str | None
    ) -> ToolCallResult:
        del data_access_token
        descriptor, data = await self.attachment_service.read_owned(
            attachment_id=self._attachment_id(request), user_token=request.user_token
        )
        self._require_spreadsheet(descriptor)
        limit = min(500, max(1, int(request.arguments.get("limit") or 100)))
        result = await asyncio.to_thread(
            self._filter_sync,
            descriptor,
            data,
            request.arguments.get("sheet_name"),
            list(request.arguments.get("filters") or []),
            [str(item) for item in request.arguments.get("select_columns") or []],
            limit,
        )
        return ToolCallResult(
            tool_id=SPREADSHEET_FILTER_TOOL_ID,
            status=ToolResultStatus.SUCCESS,
            content=result,
            structured_content=result,
            metadata={"deterministic": True},
        )

    async def spreadsheet_aggregate(
        self, request: ToolCallRequest, data_access_token: str | None
    ) -> ToolCallResult:
        del data_access_token
        descriptor, data = await self.attachment_service.read_owned(
            attachment_id=self._attachment_id(request), user_token=request.user_token
        )
        self._require_spreadsheet(descriptor)
        metrics = list(request.arguments.get("metrics") or [])
        if not metrics:
            raise AppError("AGGREGATION_METRICS_REQUIRED", "metrics 不能为空", 422)
        limit = min(500, max(1, int(request.arguments.get("limit") or 100)))
        result = await asyncio.to_thread(
            self._aggregate_sync,
            descriptor,
            data,
            request.arguments.get("sheet_name"),
            list(request.arguments.get("filters") or []),
            [str(item) for item in request.arguments.get("group_by") or []],
            metrics,
            limit,
        )
        return ToolCallResult(
            tool_id=SPREADSHEET_AGGREGATE_TOOL_ID,
            status=ToolResultStatus.SUCCESS,
            content=result,
            structured_content=result,
            metadata={"deterministic": True},
        )

    @staticmethod
    def _require_spreadsheet(descriptor: AttachmentDescriptor) -> None:
        extension = Path(descriptor.filename).suffix.lower()
        if descriptor.kind != AttachmentKind.SPREADSHEET:
            raise AppError("SPREADSHEET_REQUIRED", "该工具只能处理表格附件", 422)
        if extension not in {".xlsx", ".csv", ".tsv"}:
            raise AppError(
                "SPREADSHEET_FORMAT_NOT_SUPPORTED",
                "本地确定性表格工具仅支持 XLSX、CSV 和 TSV；其他格式可由通用内容分析智能体处理。",
                422,
            )

    @staticmethod
    def _extension(descriptor: AttachmentDescriptor) -> str:
        return Path(descriptor.filename).suffix.lower()

    def _inspect_sync(self, descriptor: AttachmentDescriptor, data: bytes) -> dict[str, Any]:
        extension = self._extension(descriptor)
        if extension == ".xlsx":
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                sheets = []
                for sheet in workbook.worksheets:
                    preview = []
                    for row in sheet.iter_rows(min_row=1, max_row=6, values_only=True):
                        preview.append([self._json_value(value) for value in row[:50]])
                    sheets.append(
                        {
                            "name": sheet.title,
                            "rows": sheet.max_row,
                            "columns": sheet.max_column,
                            "headers": preview[0] if preview else [],
                            "preview_rows": preview[1:6],
                        }
                    )
                return {"filename": descriptor.filename, "format": "xlsx", "sheets": sheets}
            finally:
                workbook.close()
        rows, delimiter, encoding = self._csv_rows(data, extension)
        preview = []
        total = 0
        max_columns = 0
        scan_truncated = False
        for row in rows:
            if total >= self.max_scan_rows:
                scan_truncated = True
                break
            total += 1
            max_columns = max(max_columns, len(row))
            if len(preview) < 6:
                preview.append(row[:50])
        return {
            "filename": descriptor.filename,
            "format": "tsv" if delimiter == "\t" else "csv",
            "encoding": encoding,
            "sheets": [
                {
                    "name": "CSV",
                    "rows_scanned": total,
                    "scan_truncated": scan_truncated,
                    "columns": max_columns,
                    "headers": preview[0] if preview else [],
                    "preview_rows": preview[1:6],
                }
            ],
        }

    def _read_range_sync(
        self,
        descriptor: AttachmentDescriptor,
        data: bytes,
        sheet_name: Any,
        min_row: int,
        max_row: int,
        min_col: int,
        max_col: int,
    ) -> dict[str, Any]:
        extension = self._extension(descriptor)
        values: list[list[Any]] = []
        selected_sheet = str(sheet_name or "")
        if extension == ".xlsx":
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                sheet = self._worksheet(workbook, selected_sheet)
                selected_sheet = sheet.title
                for row in sheet.iter_rows(
                    min_row=min_row,
                    max_row=max_row,
                    min_col=min_col,
                    max_col=max_col,
                    values_only=True,
                ):
                    values.append([self._json_value(value) for value in row])
            finally:
                workbook.close()
        else:
            rows, _, _ = self._csv_rows(data, extension)
            for index, row in enumerate(rows, start=1):
                if index < min_row:
                    continue
                if index > max_row:
                    break
                values.append(row[min_col - 1 : max_col])
            selected_sheet = "CSV"
        return {
            "filename": descriptor.filename,
            "sheet_name": selected_sheet,
            "cell_range": self._range_label(min_row, max_row, min_col, max_col),
            "rows": values,
            "row_count": len(values),
        }

    def _filter_sync(
        self,
        descriptor: AttachmentDescriptor,
        data: bytes,
        sheet_name: Any,
        filters: list[dict[str, Any]],
        select_columns: list[str],
        limit: int,
    ) -> dict[str, Any]:
        selected_sheet, headers, rows = self._tabular_rows(descriptor, data, sheet_name)
        indexes = self._header_indexes(headers)
        self._validate_columns(indexes, [str(item.get("column") or "") for item in filters])
        if select_columns:
            self._validate_columns(indexes, select_columns)
        else:
            select_columns = list(headers)
        output: list[dict[str, Any]] = []
        matched_count = 0
        scanned_count = 0
        scan_truncated = False
        for row in rows:
            if scanned_count >= self.max_scan_rows:
                scan_truncated = True
                break
            scanned_count += 1
            record = self._record(headers, row)
            if self._matches(record, filters):
                matched_count += 1
                if len(output) < limit:
                    output.append({column: record.get(column) for column in select_columns})
        return {
            "filename": descriptor.filename,
            "sheet_name": selected_sheet,
            "headers": select_columns,
            "rows": output,
            "returned_count": len(output),
            "matched_count": matched_count,
            "rows_scanned": scanned_count,
            "scan_truncated": scan_truncated,
        }

    def _aggregate_sync(
        self,
        descriptor: AttachmentDescriptor,
        data: bytes,
        sheet_name: Any,
        filters: list[dict[str, Any]],
        group_by: list[str],
        metrics: list[dict[str, Any]],
        limit: int,
    ) -> dict[str, Any]:
        selected_sheet, headers, rows = self._tabular_rows(descriptor, data, sheet_name)
        indexes = self._header_indexes(headers)
        filter_columns = [str(item.get("column") or "") for item in filters]
        metric_columns = [
            str(item.get("column") or "")
            for item in metrics
            if str(item.get("operation") or "").lower() != "count" or item.get("column")
        ]
        self._validate_columns(indexes, [*filter_columns, *group_by, *metric_columns])

        groups: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
        scanned_count = 0
        matched_count = 0
        scan_truncated = False
        for row in rows:
            if scanned_count >= self.max_scan_rows:
                scan_truncated = True
                break
            scanned_count += 1
            record = self._record(headers, row)
            if self._matches(record, filters):
                matched_count += 1
                groups[tuple(record.get(column) for column in group_by)].append(record)

        results: list[dict[str, Any]] = []
        for key, records in groups.items():
            item = {column: self._json_value(value) for column, value in zip(group_by, key, strict=True)}
            for metric in metrics:
                operation = str(metric.get("operation") or "count").lower()
                column = str(metric.get("column") or "")
                alias = str(metric.get("alias") or f"{operation}_{column or 'rows'}")
                item[alias] = self._metric(records, operation, column)
            results.append(item)

        results.sort(key=lambda item: tuple(str(item.get(column) or "") for column in group_by))
        return {
            "filename": descriptor.filename,
            "sheet_name": selected_sheet,
            "group_by": group_by,
            "metrics": metrics,
            "rows": results[:limit],
            "group_count": len(results),
            "returned_count": min(len(results), limit),
            "matched_count": matched_count,
            "rows_scanned": scanned_count,
            "scan_truncated": scan_truncated,
        }

    def _tabular_rows(
        self, descriptor: AttachmentDescriptor, data: bytes, sheet_name: Any
    ) -> tuple[str, list[str], Any]:
        extension = self._extension(descriptor)
        if extension == ".xlsx":
            # Materialize rows in the worker thread so the workbook can be closed safely.
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                sheet = self._worksheet(workbook, str(sheet_name or ""))
                iterator = sheet.iter_rows(values_only=True)
                first = next(iterator, ())
                headers = self._headers(first)
                rows = [list(row) for _, row in zip(range(self.max_scan_rows + 1), iterator, strict=False)]
                return sheet.title, headers, rows
            finally:
                workbook.close()
        rows, _, _ = self._csv_rows(data, extension)
        first = next(rows, [])
        return "CSV", self._headers(first), rows

    @staticmethod
    def _worksheet(workbook, sheet_name: str):
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise AppError(
                    "SHEET_NOT_FOUND",
                    f"工作表不存在: {sheet_name}；可用工作表: {', '.join(workbook.sheetnames)}",
                    404,
                )
            return workbook[sheet_name]
        if not workbook.worksheets:
            raise AppError("SPREADSHEET_EMPTY", "表格没有工作表", 422)
        return workbook.worksheets[0]

    @staticmethod
    def _headers(row: Any) -> list[str]:
        values: list[str] = []
        seen: dict[str, int] = {}
        for index, value in enumerate(row, start=1):
            name = str(value).strip() if value is not None else f"column_{index}"
            name = name or f"column_{index}"
            count = seen.get(name, 0) + 1
            seen[name] = count
            values.append(name if count == 1 else f"{name}_{count}")
        return values

    @staticmethod
    def _header_indexes(headers: list[str]) -> dict[str, int]:
        return {name: index for index, name in enumerate(headers)}

    @staticmethod
    def _validate_columns(indexes: dict[str, int], columns: list[str]) -> None:
        missing = sorted({column for column in columns if column and column not in indexes})
        if missing:
            raise AppError(
                "SPREADSHEET_COLUMN_NOT_FOUND",
                f"列不存在: {', '.join(missing)}；可用列: {', '.join(indexes)}",
                422,
            )

    @staticmethod
    def _record(headers: list[str], row: list[Any]) -> dict[str, Any]:
        return {
            header: FileAndSpreadsheetToolHandlers._json_value(
                row[index] if index < len(row) else None
            )
            for index, header in enumerate(headers)
        }

    @classmethod
    def _matches(cls, record: dict[str, Any], filters: list[dict[str, Any]]) -> bool:
        for condition in filters:
            column = str(condition.get("column") or "")
            operator = str(condition.get("operator") or "eq").lower()
            actual = record.get(column)
            expected = condition.get("value")
            if not cls._compare(actual, operator, expected):
                return False
        return True

    @classmethod
    def _compare(cls, actual: Any, operator: str, expected: Any) -> bool:
        if operator == "is_empty":
            return actual is None or str(actual).strip() == ""
        if operator == "not_empty":
            return not cls._compare(actual, "is_empty", expected)
        if operator == "in":
            candidates = expected if isinstance(expected, list) else [expected]
            return any(cls._compare(actual, "eq", item) for item in candidates)
        if operator in {"contains", "starts_with", "ends_with"}:
            left = "" if actual is None else str(actual).casefold()
            right = "" if expected is None else str(expected).casefold()
            if operator == "contains":
                return right in left
            if operator == "starts_with":
                return left.startswith(right)
            return left.endswith(right)
        left_num = cls._number(actual)
        right_num = cls._number(expected)
        if left_num is not None and right_num is not None:
            left: Any = left_num
            right: Any = right_num
        else:
            left = "" if actual is None else str(actual).casefold()
            right = "" if expected is None else str(expected).casefold()
        return {
            "eq": left == right,
            "ne": left != right,
            "gt": left > right,
            "gte": left >= right,
            "lt": left < right,
            "lte": left <= right,
        }.get(operator, False)

    @classmethod
    def _metric(cls, records: list[dict[str, Any]], operation: str, column: str) -> Any:
        if operation == "count":
            if not column:
                return len(records)
            return sum(1 for row in records if row.get(column) not in {None, ""})
        values = [cls._number(row.get(column)) for row in records]
        numbers = [value for value in values if value is not None]
        if not numbers:
            return None
        if operation == "sum":
            return sum(numbers)
        if operation == "avg":
            return sum(numbers) / len(numbers)
        if operation == "min":
            return min(numbers)
        if operation == "max":
            return max(numbers)
        raise AppError("AGGREGATION_OPERATION_INVALID", f"不支持的聚合操作: {operation}", 422)

    @staticmethod
    def _number(value: Any) -> float | None:
        if value is None or isinstance(value, bool):
            return None
        try:
            return float(str(value).replace(",", "").strip())
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _json_value(value: Any) -> Any:
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        if hasattr(value, "isoformat"):
            return value.isoformat()
        return str(value)

    @staticmethod
    def _csv_rows(data: bytes, extension: str):
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = data.decode("utf-8", errors="replace")
            encoding = "utf-8-replace"
        delimiter = "\t" if extension == ".tsv" else ","
        if extension != ".tsv":
            try:
                delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",\t;|").delimiter
            except csv.Error:
                pass
        return iter(csv.reader(io.StringIO(text), delimiter=delimiter)), delimiter, encoding

    @staticmethod
    def _range_label(min_row: int, max_row: int, min_col: int, max_col: int) -> str:
        from openpyxl.utils import get_column_letter

        return (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{max_row}"
        )
